"""Text extraction for various file types.

Every extractor degrades gracefully: if an optional dependency or binary is
missing, we log a warning and return an empty string rather than raising, so
ingestion never crashes on a single problematic file.
"""
from __future__ import annotations

import logging
import mimetypes
from pathlib import Path

log = logging.getLogger("doccrawler.extract")

TEXT_EXTS = {".txt", ".md", ".markdown", ".rst", ".log", ".csv"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp"}

_OCR_WARNED = False


def guess_mime_type(path: Path) -> str:
    mime, _ = mimetypes.guess_type(str(path))
    if mime:
        return mime
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        return "text/plain"
    return "application/octet-stream"


def is_supported(path: Path) -> bool:
    ext = path.suffix.lower()
    return ext in TEXT_EXTS | PDF_EXTS | DOCX_EXTS | IMAGE_EXTS


def extract_text(path: Path) -> tuple[str, bool]:
    """Returns (text, ocr_used). Never raises; logs and returns '' on failure."""
    ext = path.suffix.lower()
    try:
        if ext in TEXT_EXTS:
            return _extract_plain_text(path), False
        if ext in PDF_EXTS:
            return _extract_pdf(path), False
        if ext in DOCX_EXTS:
            return _extract_docx(path), False
        if ext in IMAGE_EXTS:
            return _extract_image_ocr(path)
    except Exception as exc:  # pragma: no cover - defensive catch-all
        log.warning("Failed to extract text from %s: %s", path, exc)
        return "", False
    log.warning("Unsupported file type for extraction: %s", path)
    return "", False


def _extract_plain_text(path: Path) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return path.read_text(encoding=enc, errors="replace")
        except Exception:
            continue
    return ""


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        log.warning("pypdf is not installed; cannot extract text from PDF %s. "
                    "Install with: pip install pypdf", path)
        return ""
    try:
        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception as exc:
                log.warning("Failed extracting a page from %s: %s", path, exc)
        return "\n".join(parts)
    except Exception as exc:
        log.warning("Failed to read PDF %s: %s", path, exc)
        return ""


def _extract_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        log.warning("python-docx is not installed; cannot extract text from %s. "
                    "Install with: pip install python-docx", path)
        return ""
    try:
        d = docx.Document(str(path))
        paras = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                paras.append(" | ".join(c.text for c in row.cells))
        return "\n".join(paras)
    except Exception as exc:
        log.warning("Failed to read docx %s: %s", path, exc)
        return ""


def _extract_image_ocr(path: Path) -> tuple[str, bool]:
    global _OCR_WARNED
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        if not _OCR_WARNED:
            log.warning("pytesseract/Pillow not available; OCR skipped for images "
                        "(e.g. %s). Install with: pip install pytesseract Pillow, "
                        "and ensure the 'tesseract' binary is installed "
                        "(Termux: pkg install tesseract).", path)
            _OCR_WARNED = True
        return "", False
    try:
        text = pytesseract.image_to_string(Image.open(path))
        return text, True
    except FileNotFoundError:
        if not _OCR_WARNED:
            log.warning("The 'tesseract' binary was not found on PATH; OCR skipped. "
                        "On Termux run: pkg install tesseract. On Debian/Ubuntu: "
                        "apt install tesseract-ocr.")
            _OCR_WARNED = True
        return "", False
    except Exception as exc:
        log.warning("OCR failed for %s: %s", path, exc)
        return "", False
